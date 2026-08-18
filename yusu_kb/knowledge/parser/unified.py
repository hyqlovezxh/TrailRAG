"""Unified document parser — local lightweight engines only.

Ported from YUSU ``yuxi.knowledge.parser.unified`` (demo edition):
- MinIO/MultiModalService/OCR/LLM-based parsing removed;
- ``.doc``, ``.xls`` and image extensions dropped from the supported set;
- pptx extracted via stdlib zipfile + XML instead of ``python-pptx``;
- csv/xlsx are emitted as markdown tables in 200-row batches.
"""

from __future__ import annotations

import asyncio
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from yusu_kb.utils.logger import logger

SUPPORTED_FILE_EXTENSIONS = {
    ".txt",
    ".md",
    ".docx",
    ".html",
    ".htm",
    ".json",
    ".csv",
    ".xlsx",
    ".pdf",
    ".pptx",
    ".zip",
}

_TABLE_BATCH_ROWS = 200


def is_supported_file_extension(filename: str) -> bool:
    """Check whether the filename's extension is supported."""
    return Path(filename).suffix.lower() in SUPPORTED_FILE_EXTENSIONS


@dataclass
class MarkdownParseResult:
    source_file_path: str
    markdown_content: str
    tables: list[str] = field(default_factory=list)
    images: list[str] = field(default_factory=list)


def _decode_text_with_fallback(raw: bytes, encodings: tuple[str, ...]) -> str:
    for encoding in encodings:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode(encodings[-1], errors="replace")


def _read_text_file(path: str) -> str:
    raw = Path(path).read_bytes()
    return _decode_text_with_fallback(raw, ("utf-8-sig", "utf-8", "gbk", "gb18030"))


def _read_csv_encoding(path: str) -> str:
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "utf-8"


def _values_to_row(values: list) -> str:
    cells = []
    for value in values:
        if value is None:
            cells.append("")
        else:
            text = str(value).replace("|", "\\|").replace("\n", " ")
            cells.append(text.strip())
    return "| " + " | ".join(cells) + " |"


def _rows_to_markdown_tables(header: list[str], rows: list[list]) -> list[str]:
    """Split a big table into markdown tables of ``_TABLE_BATCH_ROWS`` rows."""
    tables: list[str] = []
    if not header:
        return tables

    header_row = _values_to_row(header)
    separator_row = "| " + " | ".join(["---"] * len(header)) + " |"

    for start in range(0, len(rows), _TABLE_BATCH_ROWS):
        batch = rows[start : start + _TABLE_BATCH_ROWS]
        lines = [header_row, separator_row]
        lines.extend(_values_to_row(row) for row in batch)
        tables.append("\n".join(lines))
    return tables


def _parse_txt(path: str) -> str:
    return _read_text_file(path)


def _parse_markdown(path: str) -> str:
    return _read_text_file(path)


def _parse_json(path: str) -> str:
    raw = _read_text_file(path)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSON 解析失败: {exc}") from exc
    return json.dumps(data, ensure_ascii=False, indent=2)


def _parse_html(path: str) -> str:
    from markdownify import markdownify as md

    raw = _read_text_file(path)
    return md(raw, heading_style="ATX")


def _parse_pdf(path: str) -> str:
    import fitz  # PyMuPDF

    parts: list[str] = []
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ValueError(f"PDF 解析失败: {exc}") from exc
    try:
        for page in doc:
            text = page.get_text("text").strip()
            if text:
                parts.append(text)
    finally:
        doc.close()
    if not parts:
        raise ValueError("PDF 未提取到文本（可能为扫描件，暂不支持 OCR）")
    return "\n\n".join(parts)


def _parse_docx(path: str) -> str:
    import docx2txt

    text = docx2txt.process(path) or ""
    if not text.strip():
        raise ValueError("DOCX 未提取到文本")
    return text.strip()


def _parse_xlsx(path: str, table_to_markdown: bool) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(value) if value is not None else "" for value in rows[0]]
            body = [list(row) for row in rows[1:]]
            if table_to_markdown:
                tables = _rows_to_markdown_tables(header, body)
                sections.append(f"## {sheet.title}\n\n" + "\n\n".join(tables))
            else:
                lines = ["\t".join(header)]
                lines.extend("\t".join(str(value) if value is not None else "" for value in row) for row in body)
                sections.append(f"## {sheet.title}\n\n" + "\n".join(lines))
    finally:
        workbook.close()
    if not sections:
        raise ValueError("XLSX 未提取到数据")
    return "\n\n".join(sections)


def _parse_csv(path: str, table_to_markdown: bool) -> str:
    import csv as csv_module

    encoding = _read_csv_encoding(path)
    with open(path, encoding=encoding, newline="") as fh:
        reader = csv_module.reader(fh)
        rows = [row for row in reader if any(cell.strip() for cell in row)]

    if not rows:
        raise ValueError("CSV 未提取到数据")

    header = rows[0]
    body = rows[1:]
    if table_to_markdown:
        tables = _rows_to_markdown_tables(header, body)
        return "\n\n".join(tables)
    lines = ["\t".join(header)]
    lines.extend("\t".join(row) for row in body)
    return "\n".join(lines)


_SLIDE_TEXT_RE = re.compile(r"<a:t>(.*?)</a:t>", re.DOTALL)


def _parse_pptx(path: str) -> str:
    import zipfile

    slides: list[tuple[int, str]] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
            if not match:
                continue
            try:
                xml = archive.read(name).decode("utf-8", errors="replace")
            except (RuntimeError, zipfile.BadZipFile) as exc:
                logger.warning(f"pptx: failed to read {name}: {exc}")
                continue
            texts = _SLIDE_TEXT_RE.findall(xml)
            texts = [text.strip() for text in texts if text.strip()]
            if texts:
                slides.append((int(match.group(1)), "\n".join(texts)))

    slides.sort(key=lambda item: item[0])
    if not slides:
        raise ValueError("PPTX 未提取到文本")
    return "\n\n".join(f"## 幻灯片 {index}\n\n{text}" for index, text in slides)


def _parse_zip(path: str, images_dir: str | None) -> str:
    from . import zip_utils

    data = Path(path).read_bytes()

    def save_image(name: str, content: bytes) -> None:
        if not images_dir:
            return
        images_path = Path(images_dir)
        images_path.mkdir(parents=True, exist_ok=True)
        safe_name = name.replace("\\", "/").rsplit("/", 1)[-1]
        (images_path / safe_name).write_bytes(content)

    return zip_utils.parse_zip_to_markdown(data, save_image)


_ENGINE_BY_EXTENSION: dict[str, callable] = {
    ".txt": _parse_txt,
    ".md": _parse_markdown,
    ".json": _parse_json,
    ".html": _parse_html,
    ".htm": _parse_html,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".csv": _parse_csv,
    ".pptx": _parse_pptx,
    ".zip": _parse_zip,
}


class Parser:
    """Unified entry point for document parsing."""

    @staticmethod
    async def aparse(
        source: str,
        params: dict | None = None,
        *,
        images_dir: str | None = None,
    ) -> str:
        """Parse a local file to markdown.

        Args:
            source: local path of the document
            params: parse options (``table_to_markdown`` honored for csv/xlsx)
            images_dir: where zip-embedded images are written so relative
                ``images/<name>`` links in the markdown resolve

        Returns:
            markdown content
        """
        if not os.path.isfile(source):
            raise ValueError(f"源文件不存在: {source}")

        ext = Path(source).suffix.lower()
        if ext not in SUPPORTED_FILE_EXTENSIONS:
            raise ValueError(f"不支持的文档类型: {source}")

        params = params or {}
        table_to_markdown = bool(params.get("table_to_markdown", True))
        multimodal_ocr = bool(params.get("multimodal_ocr", False))
        if multimodal_ocr:
            raise ValueError("OCR 解析暂不支持（演示版仅提供轻量文本解析）")

        engine = _ENGINE_BY_EXTENSION[ext]
        if ext in {".xlsx", ".csv"}:
            engine_func = lambda p: engine(p, table_to_markdown)
        elif ext == ".zip":
            engine_func = lambda p: engine(p, images_dir)
        else:
            engine_func = engine

        content = await asyncio.to_thread(engine_func, source)
        logger.info(f"Parsed {source} ({ext}) -> {len(content)} chars")
        return content

    @staticmethod
    async def aparse_to_result(source: str, params: dict | None = None) -> MarkdownParseResult:
        """Parse and wrap the result in a ``MarkdownParseResult``."""
        markdown = await Parser.aparse(source, params)
        return MarkdownParseResult(
            source_file_path=source,
            markdown_content=markdown,
        )